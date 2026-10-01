"""Offline regressions for the audited trust failures, using synthetic content only."""
import asyncio
import copy
import io
import json
import uuid
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import AsyncMock, MagicMock

import docx
import pytest
from fastapi.testclient import TestClient

from app.main import app
from app.database.connection import get_db_connection
from app.database.editing_repo import editing_repo
from app.database.final_report_repo import final_report_repo
from app.database.proofreading_repo import proofreading_repo
from app.database.report_processing_repo import report_processing_repo
from app.database.session_repo import session_repo
from app.report_processing.engine import ReportProcessingEngine
from app.report_processing.output_schema import ReportDraftOutput
from app.services.bible_context_service import BibleContextService, bible_context_service
from app.services.gemini_gateway import GeminiUnavailableError

client = TestClient(app)


def valid_output():
    return {
        "report_title": "Synthetic review", "theme": "Source review", "scripture_reference": "",
        "minister": "Synthetic speaker", "service_date": "2026-09-30",
        "report_text": "The synthetic speaker asked the reviewers to preserve the recorded words.",
        "reporter_extraction": {"speaker": "Synthetic speaker", "theme": "Source review",
            "scriptures_cited": [], "main_divisions": ["Preservation"], "illustrations": [], "key_admonitions": []},
        "editorial_selection": {"kept_elements": ["Recorded words"], "compressed_elements": [], "omitted_elements": []},
        "writing_notes": {"structure_summary": "Brief report", "theological_focus": ""},
        "proofreading": {"scripture_checks": [], "names_checked": [], "grammar_checked": True,
            "anti_slop_passed": True, "proofreader_notes": "Synthetic checks; human review required"},
    }


async def seed_session():
    sid = 'synthetic_' + uuid.uuid4().hex
    await session_repo.create_session(sid, title='Synthetic source', raw_text='Original synthetic raw source',
                                     verified_text='The synthetic speaker asked the reviewers to preserve the recorded words.')
    return sid


def gateway(text=None, configured=True, failure=None):
    result = SimpleNamespace(model_name='fixture-model', provider_slot='fixture', response=SimpleNamespace(
        text=text if text is not None else json.dumps(valid_output()), usage_metadata=None))
    return SimpleNamespace(is_configured=lambda: configured,
                           generate=AsyncMock(return_value=result, side_effect=failure))


async def execute(sid, provider):
    run = await report_processing_repo.create_run(sid)
    await ReportProcessingEngine(gateway=provider)._execute_pipeline(run['run_id'], sid)
    return await report_processing_repo.get_run(run['run_id'])


def export_responses(sid):
    return [client.get(f'/api/final-report/sessions/{sid}/download'),
            client.post(f'/api/report-processing/generate-docx/{sid}'),
            client.get(f'/api/report-processing/download-docx/{sid}')]


@pytest.mark.asyncio
@pytest.mark.parametrize('configured,failure', [(False, None), (True, GeminiUnavailableError('Provider unavailable'))])
async def test_unavailable_provider_preserves_sources_and_previous_report(configured, failure):
    sid = await seed_session()
    previous = await final_report_repo.finalize_report(sid, None, 'Existing synthetic draft', 'Previously reviewed source text.')
    run = await execute(sid, gateway(configured=configured, failure=failure))
    assert run['status'] == 'failed' and run['error_code'] == 'provider_unavailable'
    assert run['final_report_id'] is None
    assert (await final_report_repo.get_active_final_report(sid))['id'] == previous['id']
    assert (await session_repo.get_session(sid))['raw_text'] == 'Original synthetic raw source'
    retry = await execute(sid, gateway())
    assert retry['status'] == 'completed'
    assert (await report_processing_repo.get_latest_run_for_session(sid))['run_id'] == retry['run_id']
    assert len(await final_report_repo.list_final_report_revisions(sid)) == 2


bad_outputs = ['', '{broken', '[]', '{}', 'null', '{"report_text":123}']
for change in ({'report_title': '   '}, {'report_text': ''}, {'report_text': '   '}, {'report_text': 'tapestry'},
               {'proofreading': {**valid_output()['proofreading'], 'grammar_checked': False}},
               {'proofreading': {**valid_output()['proofreading'], 'anti_slop_passed': False}}):
    bad_outputs.append(json.dumps({**valid_output(), **change}))


@pytest.mark.asyncio
@pytest.mark.parametrize('output', bad_outputs)
async def test_invalid_output_never_persists_a_report(output):
    sid = await seed_session()
    run = await execute(sid, gateway(output))
    assert run['status'] == 'failed' and run['error_code'] == 'invalid_output'
    assert not await final_report_repo.list_final_report_revisions(sid)
    assert not await editing_repo.get_active_edited_report(sid)
    assert (await session_repo.get_session(sid))['raw_text'] == 'Original synthetic raw source'


@pytest.mark.asyncio
async def test_valid_short_draft_explicit_approval_idempotency_and_exact_export():
    sid = await seed_session()
    provider = gateway()
    run = await execute(sid, provider)
    assert run['status'] == 'completed' and run['review_status'] == 'needs_review'
    assert run['draft_report_id'] == run['final_report_id']
    assert provider.generate.await_count == 1
    assert provider.generate.call_args.kwargs['config'].response_schema is ReportDraftOutput
    draft = await final_report_repo.get_active_final_report(sid)
    assert draft['source_run_id'] == run['run_id'] and draft['approval_status'] == 'draft'
    assert all(response.status_code == 409 for response in export_responses(sid))
    detail = client.get(f'/api/final-report/sessions/{sid}').json()
    assert detail['can_approve'] and not detail['can_export']
    assert not await proofreading_repo.get_active_proofread_report(sid)
    for _ in range(2):
        approved = client.post(f'/api/final-report/sessions/{sid}/approve', json={'revision_id': draft['id']})
        assert approved.status_code == 200 and approved.json()['final_report']['can_export']
    async with get_db_connection() as conn:
        cur = await conn.execute('SELECT COUNT(*) FROM report_approvals WHERE session_id=?', (sid,))
        assert (await cur.fetchone())[0] == 1
    responses = export_responses(sid)
    assert all(response.status_code == 200 for response in responses)
    for response in (responses[0], responses[2]):
        document = docx.Document(io.BytesIO(response.content))
        assert valid_output()['report_text'] in '\n'.join(p.text for p in document.paragraphs)
    assert (await session_repo.get_session(sid))['final_report_status'] == 'approved'
    settings = await report_processing_repo.get_all_settings()
    assert settings['auto_continue_to_proofreading'] == 'false'


@pytest.mark.asyncio
async def test_human_edit_invalidates_approval_and_stale_revision_cannot_be_approved():
    sid = await seed_session()
    await execute(sid, gateway())
    draft = await final_report_repo.get_active_final_report(sid)
    await final_report_repo.approve_report(sid, draft['id'])
    saved = client.post(f'/api/final-report/sessions/{sid}/save-revision', json={'report_text': 'Human revised synthetic words.'}).json()['final_report']
    assert saved['approval_status'] == 'draft' and not saved['can_export']
    assert all(response.status_code == 409 for response in export_responses(sid))
    assert client.post(f'/api/final-report/sessions/{sid}/approve', json={'revision_id': draft['id']}).status_code == 409
    assert client.post(f'/api/final-report/sessions/{sid}/approve', json={'revision_id': saved['id']}).status_code == 200
    restored = await final_report_repo.activate_final_report_revision(sid, draft['id'])
    assert restored['approval_status'] == 'draft' and not restored['can_export']
    history = await final_report_repo.list_final_report_revisions(sid)
    assert len(history) == 2 and history[1]['approved_at']


@pytest.mark.asyncio
@pytest.mark.parametrize('source', ['editor', 'proofreader', 'verified', 'raw'])
async def test_source_edits_revoke_current_export(source):
    sid = await seed_session()
    await execute(sid, gateway())
    draft = await final_report_repo.get_active_final_report(sid)
    await final_report_repo.approve_report(sid, draft['id'])
    if source == 'editor':
        await editing_repo.save_edited_report_revision(sid, 'Changed synthetic editor words.', 'human_edited', 1, 'v1')
    elif source == 'proofreader':
        await proofreading_repo.save_proofread_revision(sid, 'Changed synthetic proofreader words.', 'human_reviewed', 1, 'v1')
    else:
        async with get_db_connection() as conn:
            field = 'raw_text' if source == 'raw' else 'verified_text'
            await conn.execute(f'UPDATE sessions SET {field}=? WHERE session_id=?', ('Changed synthetic source', sid))
            await conn.commit()
    assert all(response.status_code == 409 for response in export_responses(sid))
    assert not (await final_report_repo.get_active_final_report(sid))['can_export']
    assert client.post(f'/api/final-report/sessions/{sid}/approve', json={'revision_id': draft['id']}).status_code == 409


@pytest.mark.asyncio
async def test_archive_restore_retains_original_bytes_raw_text_and_revisions(tmp_path):
    sid = await seed_session()
    source = (tmp_path / 'original.wav').resolve()
    assert tmp_path.resolve() in source.parents
    original = b'RIFF-synthetic-disposable-WAV-source'
    source.write_bytes(original)
    async with get_db_connection() as conn:
        await conn.execute('UPDATE sessions SET audio_file_path=? WHERE session_id=?', (str(source), sid))
        await conn.commit()
    await execute(sid, gateway())
    for _ in range(2):
        assert client.delete(f'/api/sessions/{sid}').json()['status'] == 'archived'
    assert source.read_bytes() == original
    assert sid not in {s['session_id'] for s in client.get('/api/sessions').json()['sessions']}
    assert sid in {s['session_id'] for s in client.get('/api/sessions?archived=true').json()['sessions']}
    assert sid in {s['session_id'] for s in client.get('/api/sessions?include_archived=true').json()['sessions']}
    assert len(await final_report_repo.list_final_report_revisions(sid)) == 1
    assert client.post('/api/report-processing/start', json={'session_id': sid}).status_code == 409
    for _ in range(2):
        assert client.post(f'/api/sessions/{sid}/restore').json()['status'] == 'restored'
    assert (await session_repo.get_session(sid))['raw_text'] == 'Original synthetic raw source'
    assert source.read_bytes() == original


def test_sibling_recording_delete_cannot_remove_original(tmp_path, monkeypatch):
    import app.audio.router as audio_router
    monkeypatch.setattr(audio_router, 'STORAGE_AUDIO_DIR', str(tmp_path))
    source = tmp_path / 'synthetic-source.wav'
    source.write_bytes(b'protected synthetic recording')
    assert client.delete('/api/audio/recordings/synthetic-source').status_code == 409
    assert source.read_bytes() == b'protected synthetic recording'


@pytest.mark.asyncio
async def test_legacy_final_migration_preserves_unreviewed_text():
    sid = await seed_session()
    async with get_db_connection() as conn:
        await conn.execute("""INSERT INTO final_reports (id, session_id, revision_number,
            report_title, report_text, is_active, created_at, updated_at)
            VALUES (?, ?, 1, 'Legacy synthetic', 'Historical synthetic text', 1, '2026-01-01', '2026-01-01')""", (sid, sid))
        await conn.execute("UPDATE sessions SET final_report_status='complete', final_report_id=? WHERE session_id=?", (sid, sid))
        await conn.commit()
    await session_repo.init_db()
    legacy = await final_report_repo.get_active_final_report(sid)
    assert legacy['approval_status'] == 'legacy_unreviewed' and legacy['approved_at'] is None
    assert legacy['report_text'] == 'Historical synthetic text'
    assert (await session_repo.get_session(sid))['final_report_status'] == 'needs_review'
    assert all(response.status_code == 409 for response in export_responses(sid))
    assert client.post(f'/api/final-report/sessions/{sid}/approve', json={'revision_id': sid}).status_code == 200


@pytest.mark.asyncio
async def test_finalize_requires_accepted_current_proofread_from_same_session():
    sid, other = await seed_session(), await seed_session()
    proofread = await proofreading_repo.save_proofread_revision(sid, 'Synthetic proofread words.', 'ai_proofread', 1, 'v1')
    assert client.post(f'/api/final-report/sessions/{sid}/finalize', json={}).status_code == 409
    await proofreading_repo.accept_revision(sid, proofread['revision_id'])
    assert client.post(f'/api/final-report/sessions/{other}/finalize', json={'proofread_revision_id': proofread['revision_id']}).status_code == 409
    first = client.post(f'/api/final-report/sessions/{sid}/finalize', json={})
    again = client.post(f'/api/final-report/sessions/{sid}/finalize', json={})
    assert first.status_code == again.status_code == 200
    assert first.json()['final_report']['id'] == again.json()['final_report']['id']
    assert first.json()['final_report']['proofread_report_revision_id'] == proofread['revision_id']
    await proofreading_repo.save_proofread_revision(sid, 'New unaccepted words.', 'human_reviewed', 1, 'v1')
    assert client.post(f'/api/final-report/sessions/{sid}/finalize', json={}).status_code == 409


def test_readiness_is_honest_with_missing_and_partial_domain_context(tmp_path):
    missing = BibleContextService(tmp_path / 'missing.sqlite').readiness()
    assert not missing['available'] and not missing['complete']
    loaded = bible_context_service.readiness()
    assert loaded['available'] and loaded['verse_corpus_complete']
    assert loaded['verse_count'] == 31102 and loaded['book_count'] == 66
    assert not loaded['proper_names_complete'] and loaded['missing_sources']
    response = client.get('/api/readiness').json()
    assert response['database_available'] and response['backend_available']
    assert not response['ai_provider']['configured'] and not response['ai_provider']['live_verified']
    assert not response['ready']


@pytest.mark.asyncio
async def test_force_retry_reuses_inflight_run():
    sid = await seed_session()
    provider = gateway()
    engine = ReportProcessingEngine(gateway=provider)
    run = await report_processing_repo.create_run(sid)
    results = await asyncio.gather(engine.start_processing(sid), engine.start_processing(sid, force_new=True))
    assert all(result['run_id'] == run['run_id'] for result in results)
    provider.generate.assert_not_called()
    await report_processing_repo.cancel_run(run['run_id'])


@pytest.mark.asyncio
async def test_duplicate_session_initialization_preserves_original_and_verified_text():
    sid = await seed_session()
    original = await session_repo.get_session(sid)
    await session_repo.create_session(sid, title='Retry with changed fields', raw_text='Replacement raw', verified_text='Replacement verified')
    after = await session_repo.get_session(sid)
    assert after['raw_text'] == original['raw_text']
    assert after['verified_text'] == original['verified_text']
    assert after['title'] == original['title']


@pytest.mark.asyncio
async def test_editor_changes_require_fresh_proofreading_and_repeat_accept_is_idempotent():
    sid = await seed_session()
    editor = await editing_repo.save_edited_report_revision(sid, 'Synthetic initial editor text.', 'human_edited', 1, 'v1')
    proofread = await proofreading_repo.save_proofread_revision(sid, 'Synthetic initial editor text.', 'ai_proofread', 1, 'v1')
    assert proofread['edited_report_revision_id'] == editor['revision_id']
    await proofreading_repo.accept_revision(sid, proofread['revision_id'])
    final = client.post(f'/api/final-report/sessions/{sid}/finalize', json={}).json()['final_report']
    await proofreading_repo.accept_revision(sid, proofread['revision_id'])
    assert (await final_report_repo.get_final_report_by_id(final['id']))['can_export']
    await editing_repo.save_edited_report_revision(sid, 'Changed synthetic editor source.', 'human_edited', 1, 'v1')
    assert client.post(f'/api/final-report/sessions/{sid}/finalize', json={}).status_code == 409
    assert client.post(f'/api/proofreading/sessions/{sid}/accept', json={'revision_id': proofread['revision_id']}).status_code == 409
    assert all(response.status_code == 409 for response in export_responses(sid))


def test_response_schema_serializes_with_installed_gemini_developer_sdk():
    from google.genai import models, types
    config = types.GenerateContentConfig(response_mime_type='application/json', response_schema=ReportDraftOutput)
    serialized = models._GenerateContentConfig_to_mldev(SimpleNamespace(vertexai=False), config, {}, {})
    assert serialized['responseMimeType'] == 'application/json'
    schema = serialized['responseSchema'].model_dump(mode='json')
    assert 'report_text' in schema['required']
    assert schema['properties']['proofreading']['properties']['grammar_checked']['type'] == 'BOOLEAN'


@pytest.mark.asyncio
@pytest.mark.parametrize('role', ['reporter', 'editor', 'proofreader'])
@pytest.mark.parametrize('bad', ['{}', '{broken', 'empty', 'cliche'])
async def test_sibling_generators_reject_invalid_provider_output(role, bad, monkeypatch):
    from google import genai
    from app.services.reporting_provider import GeminiReportingProvider
    from app.services.editing_provider import GeminiEditingProvider
    from app.services.proofreading_provider import GeminiProofreadingProvider
    if role == 'reporter':
        payload = {'report_title': 'Synthetic', 'overview': 'Synthetic overview', 'report_text': 'Synthetic report words.'}
    elif role == 'editor':
        payload = {'report_title': 'Synthetic', 'report_text': 'Synthetic report words.'}
    else:
        payload = {'proofread_title': 'Synthetic', 'proofread_text': 'Synthetic report words.'}
    field = 'proofread_text' if role == 'proofreader' else 'report_text'
    text = bad
    if bad in ('empty', 'cliche'):
        payload[field] = '' if bad == 'empty' else 'tapestry'
        text = json.dumps(payload)
    monkeypatch.setenv('GEMINI_API_KEY', 'synthetic-fixture-key')
    fake_models = SimpleNamespace(generate_content=AsyncMock(return_value=SimpleNamespace(text=text)))
    monkeypatch.setattr(genai, 'Client', lambda **kwargs: SimpleNamespace(aio=SimpleNamespace(models=fake_models)))
    standard = {'version': 1, 'version_label': 'v1'}
    if role == 'reporter':
        provider = gateway(text)
        result = await GeminiReportingProvider(gateway=provider).generate_report('reporter_a', {}, 'Synthetic verified text', standard)
        assert provider.generate.await_count == 1
    elif role == 'editor':
        result = await GeminiEditingProvider().generate_edited_report({}, 'Synthetic verified text',
            {'report_title': 'Synthetic A', 'report_text': 'Synthetic source A'},
            {'report_title': 'Synthetic B', 'report_text': 'Synthetic source B'}, standard)
        assert fake_models.generate_content.await_count == 1
    else:
        result = await GeminiProofreadingProvider().proofread_report({}, 'Synthetic edited text', 'Synthetic', standard)
        assert fake_models.generate_content.await_count == 1
    assert not result.is_success


def test_kjv_builder_preserves_sources_and_refuses_overwrite(tmp_path):
    import sys
    sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'scripts'))
    from build_kjv_database import build_database
    backend = Path(__file__).resolve().parents[1]
    corpus = backend / 'data/kjv/kjv_full.json'
    before = corpus.read_bytes()
    output = tmp_path / 'full.sqlite'
    build_database(output_path=output, corpus_path=corpus)
    ready = BibleContextService(output).readiness()
    assert ready['verse_corpus_complete'] and not ready['proper_names_complete']
    with pytest.raises(FileExistsError):
        build_database(output_path=output, corpus_path=corpus)
    assert corpus.read_bytes() == before


@pytest.mark.asyncio
async def test_indexing_does_not_replace_original_transcript_linkage(tmp_path, monkeypatch):
    import app.database.session_repo as repository_module
    monkeypatch.setattr(repository_module, 'STORAGE_TRANSCRIPTS_DIR', str(tmp_path))
    monkeypatch.setattr(repository_module, 'STORAGE_AUDIO_DIR', str(tmp_path))
    sid = await seed_session()
    recording_id = uuid.uuid4().hex
    async with get_db_connection() as conn:
        await conn.execute('UPDATE sessions SET recording_id=?, transcript_id=? WHERE session_id=?', (recording_id, 'original-transcript', sid))
        await conn.commit()
    (tmp_path / 'transcripts_manifest.json').write_text(json.dumps([{'transcript_id': 'replacement-transcript'}]))
    (tmp_path / 'replacement-transcript.json').write_text(json.dumps({'recording_id': recording_id, 'raw_text': 'Replacement raw text', 'segments': []}))
    await session_repo.index_existing_storage_files()
    saved = await session_repo.get_session(sid)
    assert saved['raw_text'] == 'Original synthetic raw source'
    assert saved['transcript_id'] == 'original-transcript'


@pytest.mark.asyncio
async def test_approved_export_uses_saved_metadata_and_correct_session_with_identical_filename():
    first, second = await seed_session(), await seed_session()
    for sid, words in [(first, 'Synthetic first approved wording.'), (second, 'Synthetic second approved wording.')]:
        draft = await final_report_repo.finalize_report(sid, None, 'Same synthetic title', words,
            minister='Approved synthetic minister', programme='Synthetic programme', service_date='2026-09-30')
        await final_report_repo.approve_report(sid, draft['id'])
    await session_repo.update_session_details(first, minister='Later metadata edit')
    for sid, words in [(first, 'Synthetic first approved wording.'), (second, 'Synthetic second approved wording.')]:
        for response in (export_responses(sid)[0], export_responses(sid)[2]):
            assert response.status_code == 200
            text = '\n'.join(p.text for p in docx.Document(io.BytesIO(response.content)).paragraphs)
            assert words in text and 'Approved synthetic minister' in text
            assert 'Later metadata edit' not in text
